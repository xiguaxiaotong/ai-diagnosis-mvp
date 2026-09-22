"""用户上传文件 → 纯文本提取（供诊断对话作为上下文）。

支持格式：
- 文本类：.txt .md .csv .json .log（UTF-8 / GBK 自动兜底）
- Word：.docx（python-docx，含表格）
- PDF：.pdf（pypdf）
- 图片：.png .jpg .jpeg .bmp .webp（RapidOCR 本地识别，无需外部 API）

限制：单文件 20MB；提取文本截断 MAX_TEXT_CHARS（防历史 token 膨胀）。
"""

import io
import json
import logging
import threading
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 20 * 1024 * 1024  # 20MB
MAX_TEXT_CHARS = 4_000

TEXT_EXTS = {".txt", ".md", ".csv", ".json", ".log"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".webp"}
SUPPORTED_EXTS = TEXT_EXTS | IMAGE_EXTS | {".docx", ".pdf"}

_FILE_TOO_LARGE = f"文件超过 {MAX_FILE_BYTES // (1024 * 1024)}MB 上限"
_OCR_LOCK = threading.Lock()
_ocr_instance = None  # RapidOCR 懒加载单例（模型加载约 1-2 秒）


def _decode_text(data: bytes) -> str:
    """UTF-8 优先，GBK 兜底（Windows 导出的 csv/log 常是 GBK）。"""
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _ocr_image(data: bytes) -> str:
    """图片 OCR（RapidOCR onnx 本地推理，中文效果好）。"""
    global _ocr_instance
    with _OCR_LOCK:
        if _ocr_instance is None:
            from rapidocr_onnxruntime import RapidOCR

            logger.info("[file_extract] 初始化 RapidOCR…")
            _ocr_instance = RapidOCR()
        result, _ = _ocr_instance(data)
    if not result:
        return ""
    # result: [[box, text, score], ...] → 按阅读顺序拼接
    return "\n".join(line[1].strip() for line in result if line[1].strip())


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(
        page.extract_text() or "" for page in reader.pages
    ).strip()


def extract_text(filename: str, data: bytes) -> str:
    """入口：文件名 + 字节 → 提取的纯文本（超长截断）。

    不支持的格式 / 空文件 / 超限 → ValueError（API 层转 422）。
    """
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTS:
        supported = "、".join(sorted(SUPPORTED_EXTS))
        raise ValueError(f"暂不支持「{ext or '无后缀'}」文件，支持：{supported}")
    if len(data) == 0:
        raise ValueError("文件内容为空")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(_FILE_TOO_LARGE)

    if ext in TEXT_EXTS:
        text = _decode_text(data)
        if ext == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=1)
            except json.JSONDecodeError:
                pass  # 非 JSON 内容按普通文本用
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".pdf":
        text = _extract_pdf(data)
    else:  # 图片
        text = _ocr_image(data)
        if not text.strip():
            raise ValueError("图片中没有识别到文字，请拍摄更清晰或改用文档上传")

    text = text.strip()
    if not text:
        raise ValueError("未能从文件中提取到任何文字内容")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + f"\n…（内容过长，仅保留前 {MAX_TEXT_CHARS} 字，如需完整分析请分段上传）"
    return text
